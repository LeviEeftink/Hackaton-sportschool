from pathlib import Path
import json
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from PIL import Image

ROOT=Path(__file__).resolve().parent
OUT=ROOT.parents[1]/'Fase 1 - aangepast.docx'
doc=Document()
sec=doc.sections[0]
sec.page_width=Inches(8.5); sec.page_height=Inches(11)
sec.top_margin=sec.bottom_margin=sec.left_margin=sec.right_margin=Inches(1)
sec.header_distance=sec.footer_distance=Inches(.492)

# Preset: standard_business_brief. Header: memo_masthead without a rule.
# Named overrides: Code 8 pt / 9 pt exact; Figure centered; Caption 9 pt;
# Table text 10 pt; title 30 pt; small source labels 9 pt.
def style(name,size,color='20394B',before=0,after=6,line=1.1,bold=False,font='Calibri'):
    s=doc.styles[name] if name in doc.styles else doc.styles.add_style(name,WD_STYLE_TYPE.PARAGRAPH)
    s.font.name=font; s.font.size=Pt(size); s.font.color.rgb=RGBColor.from_string(color); s.font.bold=bold
    f=s.paragraph_format; f.space_before=Pt(before); f.space_after=Pt(after); f.line_spacing=line
    return s
style('Normal',11,'202A33')
style('Title',30,'0B2545',0,8,1.1,True)
style('Subtitle',13,'566672',0,10)
for name,size,b,a,c in [('Heading 1',16,16,8,'2E74B5'),('Heading 2',13,12,6,'2E74B5'),('Heading 3',12,8,4,'1F4D78')]:
    style(name,size,c,b,a,1.1,True).paragraph_format.keep_with_next=True
style('Caption',9,'566672',4,6)
style('Small',9,'566672',0,5)
style('Table Text',10,'202A33',0,3)
cs=style('Code',8,'202A33',0,0,1,False,'Consolas'); cs.paragraph_format.line_spacing=Pt(9)
style('Figure',11,'202A33',0,0,1).paragraph_format.alignment=WD_ALIGN_PARAGRAPH.CENTER
for name in ('Header','Footer'): style(name,9,'657380',0,0)
sec.header.paragraphs[0].text='SPORTSCHOOL DE KAST  |  Fase 1 - Ontwerpen'
fp=sec.footer.paragraphs[0]; fp.alignment=WD_ALIGN_PARAGRAPH.RIGHT
fp.add_run('Ontwerp v1.0  |  ')
fld=OxmlElement('w:fldSimple'); fld.set(qn('w:instr'),'PAGE'); fp._p.append(fld)
doc.core_properties.title='Fase 1 - Ontwerpen - Sportschool De Kast'
doc.core_properties.subject='B1-K1-W2 - Inchecken op basis van abonnementstype'
doc.core_properties.author=''; doc.core_properties.keywords='Fase 1, klassendiagram, use case, PlantUML, B1-K1-W2'

def p(text,sty=None): return doc.add_paragraph(text,sty)
def h(text): doc.add_heading(text,1)
def sub(text): doc.add_heading(text,2)
def page(title):
    x=doc.add_heading(title,1)
    x.paragraph_format.page_break_before=True
def label(title,text):
    x=p(''); x.add_run(title+' ').bold=True; x.add_run(text)
def table(headers,rows,widths):
    t=doc.add_table(rows=1,cols=len(headers)); t.autofit=False
    pr=t._tbl.tblPr
    w=pr.find(qn('w:tblW')); w.set(qn('w:w'),'9360'); w.set(qn('w:type'),'dxa')
    ind=OxmlElement('w:tblInd'); ind.set(qn('w:w'),'120'); ind.set(qn('w:type'),'dxa'); pr.append(ind)
    layout=pr.find(qn('w:tblLayout'))
    if layout is None: layout=OxmlElement('w:tblLayout'); pr.append(layout)
    layout.set(qn('w:type'),'fixed')
    mar=OxmlElement('w:tblCellMar')
    for k,v in [('top',80),('bottom',80),('start',120),('end',120)]:
        x=OxmlElement('w:'+k); x.set(qn('w:w'),str(v)); x.set(qn('w:type'),'dxa'); mar.append(x)
    pr.append(mar)
    borders=OxmlElement('w:tblBorders')
    for k in ('top','left','bottom','right','insideH','insideV'):
        x=OxmlElement('w:'+k); x.set(qn('w:val'),'single'); x.set(qn('w:sz'),'4'); x.set(qn('w:color'),'D3DBE2'); borders.append(x)
    pr.append(borders)
    grid=t._tbl.tblGrid
    for c in list(grid): grid.remove(c)
    for width in widths:
        c=OxmlElement('w:gridCol'); c.set(qn('w:w'),str(width)); grid.append(c)
    for row in rows: t.add_row()
    for ri,values in enumerate([headers]+rows):
        row=t.rows[ri]
        trpr=row._tr.get_or_add_trPr(); ns=OxmlElement('w:cantSplit');trpr.append(ns)
        if ri==0: trpr.append(OxmlElement('w:tblHeader'))
        for ci,value in enumerate(values):
            cell=row.cells[ci]; cell.width=Inches(widths[ci]/1440)
            tcw=cell._tc.get_or_add_tcPr().find(qn('w:tcW')); tcw.set(qn('w:w'),str(widths[ci])); tcw.set(qn('w:type'),'dxa')
            cell.vertical_alignment=1
            cell.paragraphs[0].style='Table Text'; cell.paragraphs[0].add_run(str(value))
            if ri==0:
                for r in cell.paragraphs[0].runs:r.bold=True
                shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'F2F4F7');cell._tc.get_or_add_tcPr().append(shade)
    p('','Small')
    return t
def figure(name,width=None,height=None,caption=''):
    par=p('','Figure');run=par.add_run()
    pic=run.add_picture(str(ROOT/name),width=Inches(width) if width else None,height=Inches(height) if height else None)
    pic._inline.docPr.set('descr',caption)
    p(caption,'Caption')
def link(text,url,sty='Small'):
    par=p('',sty); hyp=OxmlElement('w:hyperlink');hyp.set(qn('r:id'),par.part.relate_to(url,RT.HYPERLINK,is_external=True))
    run=OxmlElement('w:r');prop=OxmlElement('w:rPr');co=OxmlElement('w:color');co.set(qn('w:val'),'2E74B5');prop.append(co);run.append(prop)
    tx=OxmlElement('w:t');tx.text=text;run.append(tx);hyp.append(run);par._p.append(hyp)
